"""Export du journal vers Excel (.xlsx) ou CSV.

Le fichier est produit directement par ``openpyxl`` : aucune installation
d'Excel n'est requise, l'export fonctionne sur un poste nu, et le résultat est
un vrai classeur avec couleurs de surlignage, filtres automatiques, volets
figés et horodatages typés « date » pour qu'Excel les trie correctement.

Les lignes sont écrites dans l'ordre exact où elles sont affichées, et seules
les lignes visibles sont exportées : le tri et les filtres en place au moment
de l'export sont donc fidèlement reproduits.

Le classeur est écrit en mode « write-only », ce qui garde une empreinte
mémoire constante même sur plusieurs centaines de milliers de lignes.
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .colors import parse_hex, to_hex
from .highlight import Highlighter, rule_style

#: Limite stricte d'une feuille Excel, en-tête comprise.
EXCEL_MAX_ROWS = 1_048_576

TIMESTAMP_NUMBER_FORMAT = "DD/MM/YYYY HH:MM:SS.000"


@dataclass
class Column:
    """Définition d'une colonne exportée."""

    key: str
    title: str
    width: float
    wrap: bool = False


COLUMNS = [
    Column("line", "N° ligne", 10),
    Column("timestamp", "Date / heure", 21),
    Column("level", "Niveau", 11),
    Column("source", "Source", 30),
    Column("code", "Code", 9),
    Column("message", "Message", 70, wrap=True),
    Column("details", "Détails", 45, wrap=True),
    Column("node_path", "Chemin du nœud", 60),
]


@dataclass
class ExportContext:
    """Informations de provenance reportées sur la feuille « Informations »."""

    host: str = ""
    ipc_name: str = ""
    project: str = ""
    source_path: str = ""
    filter_summary: str = ""
    sort_summary: str = ""


def _argb(hex_color: str) -> str:
    """Convertit ``#RRGGBB`` en ``FFRRGGBB`` attendu par openpyxl."""
    return "FF" + to_hex(parse_hex(hex_color)).lstrip("#")


def _cell_value(entry, key: str):
    if key == "line":
        return entry.index + 1
    if key == "timestamp":
        return entry.timestamp if entry.timestamp is not None else entry.timestamp_text
    if key == "message":
        return entry.message_multiline
    if key == "details":
        return entry.details_multiline
    return getattr(entry, key, "")


def export_xlsx(path: str, entries, highlighter: Highlighter,
                context: ExportContext | None = None, on_progress=None) -> str:
    """Écrit *entries* dans un classeur Excel. Renvoie le chemin du fichier.

    ``on_progress`` reçoit ``(lignes_écrites, total)`` au fil de l'écriture pour
    alimenter une barre de progression.
    """
    context = context or ExportContext()
    entries = list(entries)

    truncated = 0
    if len(entries) > EXCEL_MAX_ROWS - 1:
        truncated = len(entries) - (EXCEL_MAX_ROWS - 1)
        entries = entries[: EXCEL_MAX_ROWS - 1]

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Journal")

    # Les dimensions et les volets doivent être posés avant l'écriture des
    # lignes en mode write-only.
    for position, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = column.width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(entries) + 1}"

    header_fill = PatternFill("solid", fgColor="FF1F2933")
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    header_border = Border(bottom=Side(style="thin", color="FF7B8794"))
    header_alignment = Alignment(vertical="center", horizontal="left")

    header_cells = []
    for column in COLUMNS:
        cell = WriteOnlyCell(sheet, value=column.title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = header_alignment
        header_cells.append(cell)
    sheet.append(header_cells)

    # Les styles sont mis en cache par règle : créer un PatternFill par cellule
    # ferait exploser la taille du classeur et le temps d'écriture.
    fill_cache: dict[int, tuple[PatternFill | None, Font | None]] = {}

    def styles_for(rule_index: int):
        if rule_index not in fill_cache:
            rule = highlighter.rule_at(rule_index)
            if rule is None:
                fill_cache[rule_index] = (None, None)
            else:
                # Le classeur est lu sur fond blanc : on utilise toujours la
                # déclinaison « thème clair » de la teinte.
                style = rule_style(rule.color, dark=False)
                fill_cache[rule_index] = (
                    PatternFill("solid", fgColor=_argb(style.background)),
                    Font(color=_argb(style.foreground)),
                )
        return fill_cache[rule_index]

    top_alignment = Alignment(vertical="top")
    wrap_alignment = Alignment(vertical="top", wrap_text=True)
    total = len(entries)

    for written, entry in enumerate(entries, start=1):
        fill, font = styles_for(getattr(entry, "highlight_index", -1))
        row = []
        for column in COLUMNS:
            cell = WriteOnlyCell(sheet, value=_cell_value(entry, column.key))
            if column.key == "timestamp" and isinstance(cell.value, datetime):
                cell.number_format = TIMESTAMP_NUMBER_FORMAT
            cell.alignment = wrap_alignment if column.wrap else top_alignment
            if fill is not None:
                cell.fill = fill
                cell.font = font
            row.append(cell)
        sheet.append(row)

        if on_progress and (written % 2000 == 0 or written == total):
            on_progress(written, total)

    _append_info_sheet(workbook, context, total, truncated, highlighter)

    workbook.save(path)
    return path


def _append_info_sheet(workbook: Workbook, context: ExportContext, exported: int,
                       truncated: int, highlighter: Highlighter) -> None:
    """Feuille récapitulative : provenance, filtres appliqués, légende couleurs."""
    sheet = workbook.create_sheet("Informations")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 80

    title_font = Font(bold=True, size=13)
    label_font = Font(bold=True)

    def row(label: str = "", value="", label_style: Font | None = None):
        left = WriteOnlyCell(sheet, value=label)
        left.font = label_style or label_font
        right = WriteOnlyCell(sheet, value=value)
        right.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.append([left, right])

    row("Export du journal FT Optix", "", title_font)
    row()
    row("Date de l'export", datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
    row("IPC", context.ipc_name or "—")
    row("Adresse", context.host or "—")
    row("Projet Optix", context.project or "—")
    row("Fichier source", context.source_path or "—")
    row("Lignes exportées", exported)
    if truncated:
        row("Lignes non exportées", f"{truncated} (limite d'une feuille Excel atteinte)")
    row("Filtres appliqués", context.filter_summary or "aucun")
    row("Tri appliqué", context.sort_summary or "ordre du fichier")
    row()
    row("Légende des couleurs", "")
    for index, rule in enumerate(highlighter.active_rules):
        style = rule_style(rule.color, dark=False)
        name = WriteOnlyCell(sheet, value=rule.name or f"Règle {index + 1}")
        name.fill = PatternFill("solid", fgColor=_argb(style.background))
        name.font = Font(color=_argb(style.foreground), bold=True)
        keywords = WriteOnlyCell(sheet, value=", ".join(rule.keywords))
        keywords.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.append([name, keywords])


def export_csv(path: str, entries, on_progress=None) -> str:
    """Export CSV (séparateur ``;``, BOM UTF-8 pour qu'Excel l'ouvre correctement)."""
    entries = list(entries)
    total = len(entries)
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow([column.title for column in COLUMNS])
        for written, entry in enumerate(entries, start=1):
            values = []
            for column in COLUMNS:
                value = _cell_value(entry, column.key)
                if isinstance(value, datetime):
                    value = value.strftime("%d/%m/%Y %H:%M:%S.") + f"{value.microsecond // 1000:03d}"
                values.append(value)
            writer.writerow(values)
            if on_progress and (written % 5000 == 0 or written == total):
                on_progress(written, total)
    return path


def excel_is_available() -> bool:
    """Indique si Excel est installé (uniquement pour proposer « Ouvrir dans Excel »).

    L'export lui-même n'en dépend pas.
    """
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Excel.Application\CurVer"):
            return True
    except OSError:
        return False


def open_in_default_app(path: str) -> None:
    """Ouvre le fichier exporté avec l'application associée."""
    os.startfile(path)  # noqa: S606 - chemin produit par l'application elle-même
