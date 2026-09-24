"""Test de bout en bout du comportement en direct.

Simule un partage local contenant un journal FT Optix, puis verifie que
l'application charge le fichier, suit les ajouts, encaisse une rotation,
filtre, trie, recharge les archives et exporte de facon coherente.

Le test s'execute en rendu hors ecran : aucune fenetre n'apparait.

Lancement : .venv/Scripts/python.exe tests/test_live.py
"""
import os
import shutil
import sys
import tempfile

# Rendu hors ecran : indispensable AVANT le premier import de PySide6.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
# Le rendu hors ecran ne charge aucune police par defaut : sans cela, les
# captures produites par les tests n'affichent que des carres.
if sys.platform == "win32":
    os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", "C:/Windows"), "Fonts"))

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
OUT = tempfile.mkdtemp(prefix="ftolog_export_")

from PySide6.QtCore import QEventLoop, Qt, QTimer
from PySide6.QtWidgets import QApplication

from optixplus.common import theme
from optixplus.modules.logreader.core.config import Settings
from optixplus.modules.logreader.core.discovery import Ipc
from optixplus.common import i18n as _i18n

_i18n.install("fr")  # les vérifications portent sur les libellés français
from optixplus.modules.logreader.ui.log_tab import LogTab as MainWindow  # l'onglet reprend la fenêtre d'origine


def pump(ms):
    loop = QEventLoop()
    QTimer.singleShot(ms, loop.quit)
    loop.exec()


def line(n, level="INFO", msg="Evenement"):
    return f"07-09-2026 09:{n // 60:02d}:{n % 60:02d}.000;{level};FTOptixRuntime;;{msg} {n};;\r\n"


def main():
    app = QApplication(sys.argv)

    # On simule le partage : <racine>\Optix\Log\FTOptixRuntime.0.log
    root = tempfile.mkdtemp(prefix="ftolog_")
    log_dir = os.path.join(root, "Optix", "Log")
    os.makedirs(log_dir)
    log_path = os.path.join(log_dir, "FTOptixRuntime.0.log")
    with open(log_path, "wb") as h:
        for i in range(1, 21):
            h.write(line(i).encode())

    settings = Settings()
    settings.poll_interval_ms = 200
    settings.remember_last_host = False
    palette = theme.apply(app, "light")

    window = MainWindow(settings, palette)
    window.resize(1200, 700)
    window.show()
    pump(200)

    # On court-circuite la couche reseau : le suivi lit un chemin local.
    ipc = Ipc(host="local", netbios_name="BANC-TEST", project="ProjetSimule",
              reachable=True, share_accessible=True, log_available=True)
    from optixplus.modules.logreader.core import netshare
    real_unc = netshare.unc_path
    netshare.unc_path = lambda host, share: os.path.join(root, share)

    window.connect_to(ipc)
    pump(1200)

    ok = True

    def check(label, condition, detail=""):
        nonlocal ok
        status = "OK  " if condition else "ECHEC"
        ok = ok and condition
        print(f"  [{status}] {label}{('  -> ' + detail) if detail else ''}")

    print("=== chargement initial ===")
    check("20 lignes chargees", window.model.rowCount() == 20, f"{window.model.rowCount()}")

    def bottom_text():
        last = window.proxy.rowCount() - 1
        idx = window.proxy.index(last, 5)
        return idx.data(Qt.ItemDataRole.DisplayRole)

    def top_text():
        return window.proxy.index(0, 5).data(Qt.ItemDataRole.DisplayRole)

    check("ordre chronologique croissant (plus ancien en haut)",
          top_text() == "Evenement 1", f"haut = {top_text()!r}")
    check("la plus recente est en bas", bottom_text() == "Evenement 20",
          f"bas = {bottom_text()!r}")

    scrollbar = window.table.verticalScrollBar()
    check("vue positionnee en bas", scrollbar.value() == scrollbar.maximum(),
          f"{scrollbar.value()}/{scrollbar.maximum()}")

    print("=== ajout live de 5 lignes ===")
    with open(log_path, "ab") as h:
        for i in range(21, 26):
            h.write(line(i, "ERROR" if i % 2 else "WARNING", "Nouvel evenement").encode())
    pump(1400)

    check("25 lignes au total", window.model.rowCount() == 25, f"{window.model.rowCount()}")
    check("la nouvelle ligne est bien en bas", bottom_text() == "Nouvel evenement 25",
          f"bas = {bottom_text()!r}")
    check("defilement automatique suit le bas",
          scrollbar.value() == scrollbar.maximum(),
          f"{scrollbar.value()}/{scrollbar.maximum()}")

    print("=== rotation du fichier ===")
    shutil.move(log_path, os.path.join(log_dir, "FTOptixRuntime.1.log"))
    with open(log_path, "wb") as h:
        h.write(line(26, "INFO", "Apres rotation").encode())
    pump(1400)

    check("historique conserve apres rotation", window.model.rowCount() == 26,
          f"{window.model.rowCount()}")
    check("la ligne du nouveau fichier est en bas",
          bottom_text() == "Apres rotation 26", f"bas = {bottom_text()!r}")

    print("=== filtres ===")
    window.search_edit.setText("nouvel evenement")
    pump(300)
    check("recherche plein texte", window.proxy.rowCount() == 5,
          f"{window.proxy.rowCount()} lignes visibles")

    window.search_edit.clear()
    window._filter_on_level("ERROR")
    pump(300)
    check("filtre par niveau Erreur", window.proxy.rowCount() == 3,
          f"{window.proxy.rowCount()} lignes visibles")

    window.reset_filters()
    pump(300)
    check("reinitialisation des filtres", window.proxy.rowCount() == 26,
          f"{window.proxy.rowCount()} lignes visibles")

    print("=== tri par colonne ===")
    window.table.sortByColumn(2, Qt.SortOrder.AscendingOrder)  # Niveau
    pump(300)
    first = window.model.entry_at(
        window.proxy.mapToSource(window.proxy.index(0, 0)).row())
    check("tri par gravite : erreurs d'abord", first.level == "ERROR", first.level)
    check("resume de tri correct", "Niveau" in window.proxy.sort_summary(),
          window.proxy.sort_summary())
    window.table.sortByColumn(0, Qt.SortOrder.AscendingOrder)
    pump(200)

    print("=== chargement de l'historique archive ===")
    window.load_archives()
    pump(1500)
    check("archives ajoutees avant les lignes vives", window.model.rowCount() == 51,
          f"{window.model.rowCount()}")
    check("la plus recente reste en bas", bottom_text() == "Apres rotation 26",
          f"bas = {bottom_text()!r}")

    print("=== export du contenu visible ===")
    xlsx = os.path.join(OUT, "live_export.xlsx")
    from optixplus.modules.logreader.core import export
    from optixplus.modules.logreader.ui.log_model import COLUMNS
    entries = window.proxy.visible_entries()
    export.export_xlsx(xlsx, entries, window.highlighter)
    from openpyxl import load_workbook
    ws = load_workbook(xlsx)["Journal"]
    check("export du meme nombre de lignes", ws.max_row == 52, f"{ws.max_row - 1} lignes")

    window.close()
    pump(300)
    netshare.unc_path = real_unc
    shutil.rmtree(root, ignore_errors=True)

    print("\nRESULTAT :", "tous les controles passent" if ok else "DES CONTROLES ONT ECHOUE")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
