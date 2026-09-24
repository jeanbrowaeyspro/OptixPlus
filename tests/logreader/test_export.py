"""Export Excel du journal (cœur, sans fenêtre)."""

from __future__ import annotations

from openpyxl import load_workbook

from optixplus.modules.logreader.core import export
from optixplus.modules.logreader.core.config import default_rules
from optixplus.modules.logreader.core.highlight import Highlighter
from optixplus.modules.logreader.core.logparser import parse_line

from .conftest import line


def test_xlsx_export_writes_one_row_per_entry(tmp_path):
    highlighter = Highlighter(default_rules())
    entries = [parse_line(line(i, "ERROR" if i % 2 else "INFO").rstrip("\r\n"), i) for i in range(51)]
    highlighter.apply(entries)
    path = tmp_path / "journal.xlsx"
    progress: list[tuple[int, int]] = []
    export.export_xlsx(str(path), entries, highlighter, on_progress=lambda done, total: progress.append((done, total)))
    sheet = load_workbook(path)["Journal"]
    assert sheet.max_row == 52  # en-tête + une ligne par entrée
    assert [sheet.cell(row=2, column=c).value for c in (1, 6)] == [1, "Evenement 0"]
    assert progress[-1] == (51, 51)
